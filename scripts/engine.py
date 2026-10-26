"""Academic Report Format 1.0.0: styles, native numbering, fields and OMML."""
from pathlib import Path
from copy import deepcopy
import re, json, hashlib
from docx import Document
from docx.shared import Pt, Twips, Inches, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree as ET

ROOT=Path(__file__).resolve().parents[1]
VERSION='1.0.0'
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
M='http://schemas.openxmlformats.org/officeDocument/2006/math'
NS={'w':W,'m':M,'wp':'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'}
BODY='正文_改'
HEADINGS=['Heading 1','Heading 2','Heading 3']
ROLES={BODY:('宋体',12,False), 'Heading 1':('黑体',16,True),'Heading 2':('黑体',14,True),'Heading 3':('黑体',12,True),
 '参考文献标题':('宋体',16,True),'参考文献条目':('宋体',10.5,False),'图题':('宋体',10.5,False),'表题':('宋体',12,False),
 '图片':('宋体',12,False),'公式':('宋体',12,False),'表头':('宋体',10.5,True),'表格正文':('宋体',10.5,False),
 '封面机构':('华文行楷',24,False),'封面类型':('楷体',36,True),'封面信息':('黑体',18,False),
 'Title':('黑体',18,True),'目录标题':('宋体',15,False),'页眉页脚':('宋体',10.5,False),
 'TOC 1':('宋体',14,False),'TOC 2':('宋体',14,False),'TOC 3':('宋体',12,False)}

def el(tag,**attrs):
 e=OxmlElement(tag)
 for k,v in attrs.items():e.set(qn('w:'+k),str(v))
 return e

def replace_child(parent,tag,**attrs):
 for c in list(parent):
  if c.tag==qn(tag):parent.remove(c)
 c=el(tag,**attrs);parent.append(c);return c

def fonts(rpr,east='宋体',size=None,bold=None):
 f=replace_child(rpr,'w:rFonts',eastAsia=east,ascii='Times New Roman',hAnsi='Times New Roman',cs='Times New Roman')
 if size is not None:
  replace_child(rpr,'w:sz',val=int(size*2));replace_child(rpr,'w:szCs',val=int(size*2))
 if bold is not None:
  replace_child(rpr,'w:b',val=int(bold));replace_child(rpr,'w:bCs',val=int(bold))
 return f

def field(p,code):
 """Complex field with empty cache, never a fabricated number."""
 for typ in ['begin','instruction','separate','end']:
  r=p.add_run();node=el('w:fldChar',fldCharType=typ) if typ!='instruction' else el('w:instrText')
  if typ=='begin':node.set(qn('w:dirty'),'true')
  if typ=='instruction':node.set(qn('xml:space'),'preserve');node.text=' '+code+' '
  r._r.append(node)

def update_flag(doc):replace_child(doc.settings.element,'w:updateFields',val='true')

def configure_styles(doc):
 for name,(east,size,bold) in ROLES.items():
  s=doc.styles[name] if name in doc.styles else doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
  s.base_style=doc.styles['Normal'];s.next_paragraph_style=doc.styles[BODY] if BODY in doc.styles else s
  fonts(s.element.get_or_add_rPr(),east,size,bold)
  s.font.color.rgb=RGBColor(0,0,0)
  pf=s.paragraph_format;pf.space_before=Pt(0);pf.space_after=Pt(0)
  pf.first_line_indent=Pt(0);pf.left_indent=Pt(0);pf.right_indent=Pt(0)
  pf.line_spacing=Pt(20);pf.keep_with_next=False;pf.keep_together=False
  pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY if name in [BODY,'参考文献条目'] else WD_ALIGN_PARAGRAPH.CENTER
  replace_child(s.element.get_or_add_pPr(),'w:snapToGrid',val=0)
  for child in list(s.element.get_or_add_pPr()):
   if child.tag==qn('w:pBdr'):s.element.get_or_add_pPr().remove(child)
 s=doc.styles[BODY];replace_child(s.element.get_or_add_pPr(),'w:ind',firstLineChars=200,firstLine=480)
 for i,name in enumerate(HEADINGS):
  s=doc.styles[name];pf=s.paragraph_format
  pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 else WD_ALIGN_PARAGRAPH.LEFT
  pf.space_before=Pt([12,18,12][i]);pf.space_after=Pt([3,6,6][i]);pf.keep_with_next=True;pf.keep_together=True
  if i==2:pf.line_spacing=1.0
  replace_child(s.element.get_or_add_pPr(),'w:outlineLvl',val=i)
  pf.page_break_before=(i==0)
 r=doc.styles['参考文献条目'];r.paragraph_format.line_spacing=Pt(20);r.paragraph_format.line_spacing_rule=WD_LINE_SPACING.AT_LEAST
 replace_child(r.element.get_or_add_pPr(),'w:ind',left=525,hanging=525)
 s=doc.styles['参考文献标题'];s.paragraph_format.space_before=Pt(12);s.paragraph_format.space_after=Pt(3)
 s.paragraph_format.keep_with_next=True;s.paragraph_format.page_break_before=True
 replace_child(s.element.get_or_add_pPr(),'w:outlineLvl',val=0)
 for name in ['表题','图片','表头'] :doc.styles[name].paragraph_format.keep_with_next=True
 doc.styles['图片'].paragraph_format.line_spacing=1.0
 doc.styles['公式'].paragraph_format.line_spacing=1.0
 for i in range(1,4):
  s=doc.styles['TOC '+str(i)];s.paragraph_format.left_indent=Pt([5,10,42][i-1])
  s.paragraph_format.tab_stops.add_tab_stop(Twips(8504),2,1)
 fonts(doc.styles['Normal'].element.get_or_add_rPr(),'宋体',12,False)
 for name in ['封面机构','封面类型','Title','封面信息']:doc.styles[name].paragraph_format.line_spacing=1.5
 doc.styles['封面机构'].paragraph_format.space_after=Pt(70)
 doc.styles['封面类型'].paragraph_format.space_after=Pt(45)
 doc.styles['Title'].paragraph_format.space_after=Pt(30)
 doc.styles['封面信息'].paragraph_format.space_after=Pt(12)
 update_flag(doc)

def numbering(doc):
 root=doc.part.numbering_part.element
 aids=[int(n.get(qn('w:abstractNumId'))) for n in root.findall(qn('w:abstractNum'))]
 nids=[int(n.get(qn('w:numId'))) for n in root.findall(qn('w:num'))]
 aid=max(aids,default=-1)+1;nid=max(nids,default=49)+1
 a=el('w:abstractNum',abstractNumId=aid);a.extend([el('w:nsid',val='41525054'),el('w:multiLevelType',val='multilevel')])
 for i in range(9):
  lv=el('w:lvl',ilvl=i);lv.extend([el('w:start',val=1),el('w:numFmt',val='decimal')])
  if i<3:lv.append(el('w:pStyle',val=doc.styles[HEADINGS[i]].style_id))
  lv.extend([el('w:suff',val='space'),el('w:lvlText',val='.'.join('%'+str(k+1) for k in range(i+1))),el('w:lvlJc',val='left')])
  pp=el('w:pPr');pp.append(el('w:ind',left=0,firstLine=0));lv.append(pp)
  rp=el('w:rPr');fonts(rp,'黑体',[16,14,12][min(i,2)],True);lv.append(rp);a.append(lv)
 # abstractNum must precede num elements (CT_Numbering ordering).
 nums=root.findall(qn('w:num'));root.insert(root.index(nums[0]) if nums else len(root),a)
 n=el('w:num',numId=nid);n.append(el('w:abstractNumId',val=aid));root.append(n)
 for i,name in enumerate(HEADINGS):
  pr=replace_child(doc.styles[name].element.get_or_add_pPr(),'w:numPr');pr.extend([el('w:ilvl',val=i),el('w:numId',val=nid)])
 return nid

def page_setup(section,index):
 section.page_width=Twips(11906);section.page_height=Twips(16838)
 section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Twips(1701)
 section.header_distance=Twips(851 if index==0 else 1134);section.footer_distance=Twips(992 if index==0 else 1134)
 section.different_first_page_header_footer=False
 replace_child(section._sectPr,'w:docGrid',type='lines',linePitch=312)
 if index:
  replace_child(section._sectPr,'w:pgNumType',fmt='upperRoman' if index==1 else 'decimal',start=1)

def canonicalize(doc):
 """Order schema sequences after low-level modifications (Word round-trip stability)."""
 orders={
 'pPr':'pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange',
 'rPr':'rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs em lang eastAsianLayout specVanish oMath rPrChange',
 'sectPr':'headerReference footerReference footnotePr endnotePr type pgSz pgMar paperSrc pgBorders lnNumType pgNumType cols formProt vAlign noEndnote titlePg textDirection bidi rtlGutter docGrid printerSettings sectPrChange'}
 for root in [doc.element,doc.styles.element,doc.part.numbering_part.element]:
  for node in root.iter():
   key=ET.QName(node).localname
   if key in orders:
    ranks={k:i for i,k in enumerate(orders[key].split())}
    node[:]=sorted(node,key=lambda c:ranks.get(ET.QName(c).localname,999))

def furniture(section,index,title=''):
 for which,obj in enumerate([section.header,section.footer]):
  obj.is_linked_to_previous=False
  for c in list(obj._element):obj._element.remove(c)
  p=obj.add_paragraph(style='页眉页脚')
  if index and which==1:field(p,'PAGE')
  elif index and which==0:p.add_run('目 录' if index==1 else title)

def template(path):
 """Clean package built from measured source definitions; no source content copied."""
 d=Document();configure_styles(d)
 nr=d.part.numbering_part.element
 for c in list(nr):nr.remove(c)
 numbering(d)
 for child in list(d.element.body):
  if child.tag!=qn('w:sectPr'):d.element.body.remove(child)
 for key in ['author','last_modified_by','title','subject','comments','keywords','category','identifier','version']:setattr(d.core_properties,key,'')
 page_setup(d.sections[0],0);furniture(d.sections[0],0)
 d.add_paragraph('',style='Title');canonicalize(d);d.save(path)
 # python-docx's bundled blank package contains unrelated sample customXml/thumbnail.
 from zipfile import ZipFile,ZIP_DEFLATED
 with ZipFile(path) as z:parts={n:z.read(n) for n in z.namelist() if not n.startswith('customXml/') and n not in ['docProps/thumbnail.jpeg','word/stylesWithEffects.xml']}
 for name in list(parts):
  if name.endswith('.rels'):
   r=ET.fromstring(parts[name])
   for n in list(r):
    if any(x in n.get('Type','') for x in ['/customXml','/thumbnail','/stylesWithEffects']):r.remove(n)
   parts[name]=ET.tostring(r,xml_declaration=True,encoding='UTF-8',standalone=True)
 r=ET.fromstring(parts['[Content_Types].xml'])
 for n in list(r):
  if n.get('PartName','').startswith('/customXml/') or n.get('PartName')=='/word/stylesWithEffects.xml':r.remove(n)
 parts['[Content_Types].xml']=ET.tostring(r,xml_declaration=True,encoding='UTF-8',standalone=True)
 with ZipFile(path,'w',ZIP_DEFLATED) as z:
  for name,b in parts.items():z.writestr(name,b)

class FieldManager:
 def __init__(self,doc):self.doc=doc;self.ids={};self.counter=0
 def register(self,key,kind):
  if not isinstance(key,str) or not key:raise ValueError('Nonempty id required')
  if key in self.ids:raise ValueError('Duplicate identifier: '+key)
  name='AR_'+hashlib.sha256(key.encode()).hexdigest()[:24];self.ids[key]=(name,kind)
 def bookmark(self,p,key,writer):
  self.counter+=1;name=self.ids[key][0];p._p.append(el('w:bookmarkStart',id=self.counter,name=name));writer();p._p.append(el('w:bookmarkEnd',id=self.counter))
 def reference(self,p,key):
  if key not in self.ids:raise ValueError('Unknown reference: '+key)
  field(p,'REF '+self.ids[key][0]+' \\h')
 def number(self,p,key,kind):
  def writer():
   if kind!='Reference':
    field(p,'STYLEREF 1 \\s');p.add_run('.' if kind=='Equation' else '-')
   field(p,'SEQ '+kind+(' \\s 1' if kind!='Reference' else ''))
  self.bookmark(p,key,writer)
 def text(self,p,text):
  for token in re.split(r'(\[\[(?:ref|cite):[^\]]+\]\])',text):
   m=re.fullmatch(r'\[\[(ref|cite):([^\]]+)\]\]',token)
   if m:
    if m[1]=='cite':p.add_run('[')
    self.reference(p,m[2])
    if m[1]=='cite':p.add_run(']')
   elif token:p.add_run(token)

def math_node(spec):
 """Small explicit math AST, every result is an editable OMML structure."""
 def me(tag,*children,**attrs):
  e=ET.Element('{'+M+'}'+tag)
  for k,v in attrs.items():e.set('{'+M+'}'+k,str(v))
  for c in children:e.append(c)
  return e
 def box(tag,x):return me(tag,math_node(x))
 if isinstance(spec,(str,int,float)):
  t=me('t');t.text=str(spec);return me('r',t)
 if not isinstance(spec,dict):raise ValueError('Math must be text leaf or object')
 kind=spec['type']
 if kind=='row':return me('box',me('e',*[math_node(x) for x in spec['items']]))
 if kind=='fraction':return me('f',box('num',spec['numerator']),box('den',spec['denominator']))
 if kind in ['sub','sup']:return me('sSub' if kind=='sub' else 'sSup',box('e',spec['base']),box('sub' if kind=='sub' else 'sup',spec['script']))
 if kind=='subsup':return me('sSubSup',box('e',spec['base']),box('sub',spec['sub']),box('sup',spec['sup']))
 if kind=='sqrt':return me('rad',me('radPr',me('degHide',val='1')),me('deg'),box('e',spec['body']))
 if kind=='root':return me('rad',box('deg',spec['degree']),box('e',spec['body']))
 if kind in ['integral','sum']:
  return me('nary',me('naryPr',me('chr',val='∫' if kind=='integral' else '∑'),me('limLoc',val='undOvr')),box('sub',spec.get('lower','')),box('sup',spec.get('upper','')),box('e',spec['body']))
 if kind=='matrix':
  rows=spec['rows']
  if not rows or len({len(r) for r in rows})!=1:raise ValueError('Matrix must be rectangular')
  return me('m',*[me('mr',*[box('e',x) for x in row]) for row in rows])
 raise ValueError('Unsupported math structure: '+kind)

def equation(spec):
 if 'mathml' in spec:
  path=Path(spec.get('xsl','C:/Program Files/Microsoft Office/root/Office16/MML2OMML.XSL'))
  if not path.exists():raise ValueError('MathML requires local Microsoft MML2OMML.XSL; use AST instead')
  parser=ET.XMLParser(resolve_entities=False,no_network=True)
  result=ET.XSLT(ET.parse(str(path)))(ET.fromstring(spec['mathml'].encode(),parser))
  root=result.getroot()
  if root.tag=='{'+M+'}oMathPara':root=root.find('{'+M+'}oMath')
  if root is None or root.tag!='{'+M+'}oMath':raise ValueError('Invalid OMML conversion')
  return deepcopy(root)
 root=ET.Element('{'+M+'}oMath',nsmap={'m':M});root.append(math_node(spec['math']));return root

def table(doc,headers,rows):
 if not headers or any(len(r)!=len(headers) for r in rows):raise ValueError('Table requires rectangular rows')
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
 width=8504//len(headers)
 for c in t.columns:c.width=Twips(width)
 replace_child(t._tbl.tblPr,'w:tblW',w=8504,type='dxa')
 bd=replace_child(t._tbl.tblPr,'w:tblBorders')
 for edge in ['top','bottom','left','right','insideH','insideV']:bd.append(el('w:'+edge,val='single' if edge in ['top','bottom'] else 'nil',sz=12,color='000000'))
 for i,values in enumerate([headers]+rows):
  row=t.rows[0] if i==0 else t.add_row()
  if i==0:row._tr.get_or_add_trPr().append(el('w:tblHeader'))
  row._tr.get_or_add_trPr().append(el('w:cantSplit'))
  for c,value in zip(row.cells,values):
   c.width=Twips(width);p=c.paragraphs[0];p.style='表头' if i==0 else '表格正文';p.add_run(str(value))
   if i==0:replace_child(c._tc.get_or_add_tcPr(),'w:tcBorders').append(el('w:bottom',val='single',sz=6,color='000000'))
 return t

def markdown(text):
 blocks=[];buf=[]
 def flush():
  if buf:blocks.append({'type':'paragraph','text':' '.join(buf)});buf.clear()
 for line in text.splitlines():
  m=re.match(r'^(#{1,3})\s+(.+)',line)
  if m:flush();blocks.append({'type':'heading','level':len(m[1]),'text':m[2]})
  elif line.strip():buf.append(line.strip())
  else:flush()
 flush();return blocks

def generate(data,output,base=None):
 base=Path(base or '.');blocks=data.get('blocks',[])+markdown(data.get('markdown',''))
 if not data.get('title'):raise ValueError('Report title is required')
 d=Document(ROOT/'assets/academic-report-template.docx')
 for c in list(d.element.body):
  if c.tag!=qn('w:sectPr'):d.element.body.remove(c)
 fm=FieldManager(d)
 for b in blocks:
  if b['type'] in ['figure','table','equation']:fm.register(b['id'],{'figure':'Figure','table':'Table','equation':'Equation'}[b['type']])
 for r in data.get('references',[]):fm.register(r['id'],'Reference')
 for name,text in [('封面机构',data.get('institution','')),('封面类型',data.get('report_type','学术报告')),('Title',data['title'])]:d.add_paragraph(text,style=name)
 for k,v in data.get('metadata',{}).items():d.add_paragraph(k+'  '+str(v),style='封面信息')
 sec=d.add_section(WD_SECTION_START.NEW_PAGE);page_setup(sec,1);furniture(sec,1)
 d.add_paragraph('目 录',style='目录标题');p=d.add_paragraph(style='TOC 1');field(p,'TOC \\o "1-3" \\h \\z \\u')
 sec=d.add_section(WD_SECTION_START.NEW_PAGE);page_setup(sec,2);furniture(sec,2,data.get('header',data['title']))
 chapter=False;last_level=0
 for b in blocks:
  kind=b['type']
  if kind=='heading':
   level=b['level'];text=b['text']
   if level not in [1,2,3] or level>last_level+1:raise ValueError('Heading level jump: '+text)
   if re.match(r'^\s*\d+(?:\.\d+)*[.、\s]',text):raise ValueError('Heading text must not contain numbering: '+text)
   p=d.add_paragraph(text,style=HEADINGS[level-1]);last_level=level
   # Explicit paragraph level stabilizes Word's import/normalization of linked built-in headings.
   sid=d.styles[HEADINGS[level-1]].element
   p._p.get_or_add_pPr().append(deepcopy(sid.find('w:pPr/w:numPr',NS)))
   if level==1:chapter=True
  elif kind=='paragraph':
   p=d.add_paragraph(style=BODY)
   if 'segments' in b:
    for s in b['segments']:
     if isinstance(s,str):fm.text(p,s)
     else:p._p.append(equation(s))
   else:fm.text(p,b['text'])
  elif kind in ['figure','table','equation']:
   if not chapter:raise ValueError(kind+' must follow a chapter')
   if kind=='figure':
    p=d.add_paragraph(style='图片');path=Path(b['path']);path=path if path.is_absolute() else base/path
    from PIL import Image
    with Image.open(path) as im:ratio=im.height/im.width
    width=min(float(b.get('width_inches',5.5)),8504/1440,7.5/ratio)
    p.add_run().add_picture(str(path),width=Inches(width))
    p=d.add_paragraph(style='图题');p.add_run('图 ');fm.number(p,b['id'],'Figure');p.add_run(' '+b['caption'])
   elif kind=='table':
    p=d.add_paragraph(style='表题');p.add_run('表 ');fm.number(p,b['id'],'Table');p.add_run(' '+b['caption']);table(d,b['headers'],b['rows'])
   else:
    p=d.add_paragraph(style='公式');pf=p.paragraph_format
    pf.tab_stops.add_tab_stop(Twips(4252),1);pf.tab_stops.add_tab_stop(Twips(8504),2);pf.alignment=WD_ALIGN_PARAGRAPH.LEFT
    p.add_run('\t')
    if b.get('mode','auto')=='mathtype':p.add_run('【在此插入 MathType 公式】')
    else:p._p.append(equation(b))
    p.add_run('\t(');fm.number(p,b['id'],'Equation');p.add_run(')')
  else:raise ValueError('Unsupported block: '+kind)
 if data.get('references'):
  d.add_paragraph('参考文献',style='参考文献标题')
  for r in data['references']:
   if re.match(r'^\s*\[\d+\]',r['text']):raise ValueError('Reference text must omit manual number')
   p=d.add_paragraph(style='参考文献条目');p.add_run('[');fm.number(p,r['id'],'Reference');p.add_run('] '+r['text'])
 update_flag(d);canonicalize(d);output=Path(output);output.parent.mkdir(parents=True,exist_ok=True);d.save(output);return output
